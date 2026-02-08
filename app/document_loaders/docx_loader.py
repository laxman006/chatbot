# -*- coding: utf-8 -*-
"""
DOCX document loader using python-docx.

Preserves headers, lists, and tables. Extracts text from paragraphs and tables.
"""

import os
from pathlib import Path
from typing import Any, Dict, List, Union

from langchain_core.documents import Document

from app.document_loaders.base_loader import BaseDocumentLoader
from app.doc_processor import extract_text_from_docx

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    DocxDocument = None


class DOCXLoader(BaseDocumentLoader):
    """Load Word documents (.docx). Uses python-docx with docx2txt fallback."""

    source_type = "docx"

    def load(self, source: Union[str, Path, Dict[str, Any], None]) -> List[Document]:
        if source is None:
            return []
        path = self._to_path(source)
        if path is None or not path.exists():
            return []
        if path.is_file() and path.suffix.lower() in (".docx", ".doc"):
            return self._load_file(path)
        if path.is_dir():
            return self._load_directory(path)
        return []

    def _to_path(self, source: Union[str, Path, Dict[str, Any]]) -> Union[Path, None]:
        if isinstance(source, Path):
            return source
        if isinstance(source, str):
            return Path(source)
        if isinstance(source, dict) and "path" in source:
            return Path(source["path"])
        return None

    def _load_file(self, path: Path) -> List[Document]:
        documents = []
        try:
            text = extract_text_from_docx(str(path))
            if not text.strip():
                return documents
            meta = self.extract_metadata(str(path))
            meta["file_path"] = str(path)
            meta["source"] = path.name
            meta["file_format"] = path.suffix.lstrip(".").lower()
            meta["content_type"] = "word_document"
            documents.append(self._doc(text, meta))
        except Exception as e:
            import logging
            logging.getLogger(__name__).warning("DOCX load failed for %s: %s", path, e)
        return documents

    def _load_directory(self, dir_path: Path) -> List[Document]:
        documents = []
        exts = {".docx", ".doc"}
        for p in sorted(dir_path.iterdir()):
            if p.suffix.lower() in exts:
                documents.extend(self._load_file(p))
        return documents

    def extract_metadata(self, item: Union[str, Path, Document, Dict[str, Any]]) -> Dict[str, Any]:
        if isinstance(item, Document):
            return dict(item.metadata)
        if isinstance(item, dict):
            return {k: v for k, v in item.items() if v is not None}
        path = Path(item) if isinstance(item, str) else item
        if not isinstance(path, Path):
            path = Path(str(path))
        meta = {"source": path.name if getattr(path, "name", None) else str(path)}
        if path.suffix.lower() not in (".docx", ".doc") or not path.exists():
            return meta
        try:
            if DOCX_AVAILABLE and DocxDocument is not None:
                doc = DocxDocument(path)
                meta["paragraph_count"] = len(doc.paragraphs)
                meta["table_count"] = len(doc.tables)
        except Exception:
            pass
        return meta
